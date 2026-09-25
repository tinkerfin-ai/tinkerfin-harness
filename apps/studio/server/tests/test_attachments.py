"""附件原件保留、授权读取、文件生成和删除边界"""

import base64
import io
import json
import struct
from datetime import UTC, datetime, timedelta
from types import SimpleNamespace
from unittest.mock import AsyncMock
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from attachment_fakes import MemoryAttachmentStorage
from PIL import Image
from pydantic import BaseModel
from sqlalchemy import select

from tinkerfin_studio.api.errors import BusinessException
from tinkerfin_studio.attachments.entity import AttachmentFile
from tinkerfin_studio.attachments.processing import validate_file
from tinkerfin_studio.attachments.service import AttachmentService, byte_chunks
from tinkerfin_studio.conversation.models import ConversationThread


def png():
    output = io.BytesIO()
    Image.new("RGB", (32, 24), "blue").save(output, "PNG")
    return output.getvalue()


def pptx():
    return office_container(
        "ppt/presentation.xml",
        b'<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" />',
    )


def office_container(member: str, content: bytes) -> bytes:
    output = io.BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr(member, content)
    return output.getvalue()


@pytest.fixture
async def attachments(database, attachment_storage):
    return AttachmentService(database, attachment_storage)


async def test_upload_keeps_original_and_rejects_other_users(attachments):
    data = png()
    file = await attachments.upload(
        user_id=1, name="截图.png", chunks=byte_chunks(data)
    )
    _, original = await attachments.read(file.id, user_id=1)
    assert original == data
    _, preview = await attachments.read(file.id, user_id=1, variant="preview")
    assert Image.open(io.BytesIO(preview)).format == "JPEG"
    with pytest.raises(BusinessException) as denied:
        await attachments.read(file.id, user_id=2)
    assert denied.value.error_code.http_status == 404


@pytest.mark.parametrize("name,content", [("chart.png", png), ("slides.pptx", pptx)])
async def test_model_content_uses_image_variant_or_original_file(
    attachments, name, content
):
    """授权原生输入使用图片模型变体，文档则保留原件格式和字节"""
    data = content()
    file = await attachments.upload(user_id=1, name=name, chunks=byte_chunks(data))
    resolved = await attachments.read_content(file, user_id=1)
    if file.mime_type.startswith("image/"):
        _, expected = await attachments.read(file.id, user_id=1, variant="model")
        assert resolved.data == expected
        assert resolved.mime_type == "image/jpeg"
        with Image.open(io.BytesIO(resolved.data)) as image:
            assert image.format == "JPEG"
    else:
        assert resolved.data == data
        assert resolved.mime_type == file.mime_type
    assert (await attachments.read(file.id, user_id=1))[1] == data


@pytest.mark.parametrize("user_id,collection_id", [(2, "files"), (1, "other")])
async def test_model_content_authorizes_before_reading_storage(
    attachments, attachment_storage, monkeypatch, user_id, collection_id
):
    """跨用户或集合读取在访问附件字节前被拒绝"""
    await attachments.create_collection(
        user_id=1,
        collection_id="files",
        purpose="execution",
        attachment_ids=(),
        configuration={},
    )
    file = await attachments.upload(
        user_id=1,
        collection_id="files",
        name="note.md",
        chunks=byte_chunks(b"# private"),
    )
    read = AsyncMock(wraps=attachment_storage.read)
    monkeypatch.setattr(attachment_storage, "read", read)
    with pytest.raises(FileNotFoundError, match="附件内容不可用"):
        await attachments.read_content(
            file, user_id=user_id, collection_id=collection_id
        )
    read.assert_not_awaited()


async def test_binding_prevents_cross_thread_reuse_and_draft_deletion(
    attachments, database
):
    file = await attachments.upload(
        user_id=1, name="chart.png", chunks=byte_chunks(png())
    )
    async with database.session() as session:
        session.add(
            ConversationThread(
                user_id=1,
                thread_id="thread-a",
                title="附件测试",
                created_at=datetime.now(UTC).replace(tzinfo=None),
                updated_at=datetime.now(UTC).replace(tzinfo=None),
            )
        )
        await session.commit()
        await attachments.bind(
            session, [file.id], user_id=1, thread_id="thread-a", message_id="message-a"
        )
        await session.commit()
    with pytest.raises(BusinessException):
        await attachments.get(file.id, user_id=1, thread_id="thread-b")
    with pytest.raises(FileNotFoundError):
        await attachments.read_content(file, user_id=1, thread_id="thread-b")
    assert (
        await attachments.read_content(file, user_id=1, thread_id="thread-a")
    ).mime_type == "image/jpeg"
    with pytest.raises(BusinessException) as rejected:
        await attachments.remove_draft(file.id, user_id=1)
    assert rejected.value.error_code.http_status == 409
    assert (await attachments.list_thread(user_id=1, thread_id="thread-a"))[
        0
    ].id == file.id


async def test_cleanup_removes_expired_drafts(attachments, database, monkeypatch):
    from tinkerfin_studio.attachments import service as attachment_service

    now = datetime(2026, 9, 26, tzinfo=UTC)
    monkeypatch.setattr(
        attachment_service, "datetime", SimpleNamespace(now=lambda _tz: now)
    )
    file = await attachments.upload(
        user_id=1, name="chart.png", chunks=byte_chunks(png())
    )
    async with database.session() as session:
        row = await session.get(AttachmentFile, file.id)
        row.created_at = now.replace(tzinfo=None) - timedelta(days=2)
        await session.commit()
    await attachments.cleanup()
    with pytest.raises(BusinessException):
        await attachments.get(file.id, user_id=1)
    async with database.session() as session:
        assert await session.scalar(select(AttachmentFile.id)) is None


async def test_invalid_type_and_path_do_not_publish_files(attachments):
    for name, data in [
        ("../x.png", png()),
        ("fake.png", b"not an image"),
        ("file.exe", b"binary"),
    ]:
        with pytest.raises(BusinessException):
            await attachments.upload(user_id=1, name=name, chunks=byte_chunks(data))


async def test_pptx_upload_preserves_original_container_and_format(attachments):
    """PPTX 上传保留原件字节和 MIME 类型，可按附件身份取回"""
    data = pptx()
    file = await attachments.upload(
        user_id=1, name="门店月报.pptx", chunks=byte_chunks(data)
    )
    assert file.mime_type == (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
    _, downloaded = await attachments.read(file.id, user_id=1)
    assert downloaded == data
    assert file.size_bytes == len(data)


def test_pptx_validation_rejects_bad_encrypted_and_oversized_containers():
    """PPTX 校验拒绝坏包、加密包和解压后超限的压缩包"""
    assert validate_file("report.pptx", pptx()) == (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
    with pytest.raises(ValueError, match="容器损坏"):
        validate_file("report.pptx", b"not a zip")
    with pytest.raises(ValueError, match="文档损坏或已加密"):
        validate_file("report.pptx", office_container("ppt/slide1.xml", b"<slide />"))

    encrypted = bytearray(office_container("ppt/presentation.xml", b"<ppt />"))
    encrypted[6:8] = struct.pack("<H", 1)
    central_header = encrypted.find(b"PK\x01\x02")
    assert central_header >= 0
    encrypted[central_header + 8 : central_header + 10] = struct.pack("<H", 1)
    with pytest.raises(ValueError, match="文档损坏或已加密"):
        validate_file("report.pptx", bytes(encrypted))

    with pytest.raises(ValueError, match="解压后过大"):
        validate_file(
            "report.pptx",
            office_container("ppt/presentation.xml", b"0" * (50 * 1024 * 1024 + 1)),
        )


async def test_generated_workbook_and_pdf_have_valid_content(attachments):
    """生成的工作簿保留单元格，PDF 满足附件格式校验"""
    from openpyxl import load_workbook

    workbook = await attachments.documents.run(
        {
            "operation": "generate",
            "kind": "xlsx",
            "rows": [["quarter", "revenue"], ["Q3", "128"]],
        }
    )
    parsed = load_workbook(io.BytesIO(base64.b64decode(workbook["data"])))
    try:
        assert parsed.active is not None
        assert list(parsed.active.values) == [("quarter", "revenue"), ("Q3", "128")]
    finally:
        parsed.close()
    pdf = await attachments.documents.run(
        {"operation": "generate", "kind": "pdf", "text": "第三季度营收：128 万元"}
    )
    assert (
        validate_file("report.pdf", base64.b64decode(pdf["data"])) == "application/pdf"
    )


async def test_cancelled_storage_write_removes_staging_record_and_bytes(
    database, tmp_path
):
    """取消不得留下可见附件或无法按记录清理的文件"""
    import asyncio

    class InterruptedStorage(MemoryAttachmentStorage):
        async def put(self, key, chunks):
            async def interrupted():
                async for chunk in chunks:
                    yield chunk
                    raise asyncio.CancelledError()

            await super().put(key, interrupted())

    storage = InterruptedStorage()
    service = AttachmentService(database, storage)
    with pytest.raises(asyncio.CancelledError):
        await service.upload(user_id=1, name="cancelled.png", chunks=byte_chunks(png()))
    async with database.session() as session:
        assert await session.scalar(select(AttachmentFile.id)) is None
    assert not storage.objects


async def test_workbook_preserves_numbers_and_treats_formula_like_text_as_text(
    attachments,
):
    """生成的数据表保留数值类型，不将外部文字变成可执行公式"""
    from openpyxl import load_workbook

    result = await attachments.documents.run(
        {"operation": "generate", "kind": "xlsx", "rows": [[128, "=SUM(A1:A2)"]]}
    )
    workbook = load_workbook(io.BytesIO(base64.b64decode(result["data"])))
    sheet = workbook.active
    assert sheet is not None
    assert sheet["A1"].value == 128
    assert sheet["B1"].data_type == "s"
    workbook.close()


async def test_same_name_attachments_remain_distinct_after_service_restart(
    attachments, database, attachment_storage
):
    """同名报告按 ID 区分，重新创建服务后原件和会话引用仍可读取"""
    files = []
    originals = []
    for marker in ("报告 A：收入 128", "报告 B：收入 256"):
        data = office_container(
            "word/document.xml",
            (
                '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                f"<w:body><w:p><w:r><w:t>{marker}</w:t></w:r></w:p></w:body></w:document>"
            ).encode(),
        )
        originals.append(data)
        files.append(
            await attachments.upload(
                user_id=1, name="report.docx", chunks=byte_chunks(data)
            )
        )
    assert files[0].id != files[1].id
    async with database.session() as session:
        now = datetime.now(UTC).replace(tzinfo=None)
        session.add(
            ConversationThread(
                user_id=1,
                thread_id="reports",
                title="比较报告",
                created_at=now,
                updated_at=now,
            )
        )
        await session.commit()
        await attachments.bind(
            session,
            [file.id for file in files],
            user_id=1,
            thread_id="reports",
            message_id="question",
        )
        await session.commit()
    restored = AttachmentService(database, attachment_storage)
    assert len(await restored.list_thread(user_id=1, thread_id="reports")) == 2
    results = []
    for file in files:
        _, data = await restored.read(file.id, user_id=1, thread_id="reports")
        results.append(data)
    assert results == originals


@pytest.mark.parametrize(
    "format_name, mime",
    [("PNG", "image/png"), ("JPEG", "image/jpeg"), ("WEBP", "image/webp")],
)
async def test_generated_image_tool_preserves_actual_format_and_typed_result(
    attachments, database, monkeypatch, format_name, mime, work_file_runtime
):
    """生成工作图片不发布附件，显式交付后保留实际格式和原始字节"""
    import json

    from langchain_core.messages import ToolMessage

    from tinkerfin_agui_adapter import DeepAgentAgUiAdapter, RunIdentity
    from tinkerfin_studio.attachments import tools as media_tools
    from tinkerfin_studio.attachments.workspace_tools import (
        build_sandbox_attachment_tools,
    )

    runtime, _, files = work_file_runtime

    output = io.BytesIO()
    Image.new("RGB", (32, 24), "blue").save(output, format_name)
    data = output.getvalue()

    async def generate(model, prompt, *, allowed_origins):
        assert allowed_origins == ("http://localhost:11434",)
        return data

    monkeypatch.setattr(media_tools, "generate_image_bytes", generate)
    async with database.session() as session:
        now = datetime.now(UTC).replace(tzinfo=None)
        session.add(
            ConversationThread(
                user_id=1,
                thread_id="generated",
                title="生成图片",
                created_at=now,
                updated_at=now,
            )
        )
        await session.commit()
    tools = media_tools.build_attachment_tools(
        service=attachments,
        processor=attachments.documents,
        user_id=1,
        thread_id="generated",
        image_model=None,
        model_allowed_origins=("http://localhost:11434",),
    )
    generator = next(item for item in tools if item.name == "generate_image")
    result = await generator.ainvoke(
        {
            "type": "tool_call",
            "id": "generate",
            "name": "generate_image",
            "args": {"prompt": "blue square", "runtime": runtime},
        }
    )
    assert isinstance(result, ToolMessage)
    assert isinstance(result.content, str)
    work_file = json.loads(result.content)
    assert work_file["mime_type"] == mime
    assert files[work_file["file_path"]] == data
    assert await attachments.list_thread(user_id=1, thread_id="generated") == []
    adapter = DeepAgentAgUiAdapter(
        identity=RunIdentity(namespace="test", thread_id="generated", run_id="run")
    )
    events = adapter.process(
        {
            "type": "messages",
            "ns": (),
            "data": (result, {"lc_agent_name": None, "langgraph_node": "tools"}),
        }
    )
    wire = next(
        item.model_dump(mode="json", by_alias=True)
        for item in events
        if item.type == "TOOL_CALL_RESULT"
    )
    assert json.loads(wire["content"]) == work_file
    assert wire["attachments"] == []
    deliver = build_sandbox_attachment_tools(
        service=attachments, user_id=1, thread_id="generated"
    )[0]
    delivered = await deliver.ainvoke(
        {
            "type": "tool_call",
            "id": "deliver",
            "name": "deliver_file",
            "args": {
                "runtime": runtime,
                "file_path": work_file["file_path"],
                "name": work_file["name"],
            },
        }
    )
    events = adapter.process(
        {
            "type": "messages",
            "ns": (),
            "data": (delivered, {"lc_agent_name": None, "langgraph_node": "tools"}),
        }
    )
    wire = next(
        item.model_dump(mode="json", by_alias=True)
        for item in events
        if item.type == "TOOL_CALL_RESULT"
    )
    assert wire["content"] == ""
    assert len(wire["attachments"]) == 1
    attachment = wire["attachments"][0]
    assert attachment["mime_type"] == mime
    _, original = await attachments.read(
        attachment["id"], user_id=1, thread_id="generated"
    )
    assert original == data


@pytest.mark.parametrize(
    "payload",
    [
        {"operation": "unknown", "kind": "png", "values": [1]},
        {"operation": "generate", "kind": "xlsx", "rows": [[{"unexpected": "object"}]]},
        {
            "operation": "preview_image",
            "data": "invalid-base64",
            "max_bytes": 4096,
        },
    ],
)
async def test_file_worker_rejects_invalid_input(attachments, payload):
    with pytest.raises(ValueError):
        await attachments.documents.run(payload)


@pytest.mark.parametrize(
    "failure_type,kind",
    [(ValueError, "invalid_input"), (RuntimeError, "internal_error")],
)
def test_file_worker_reports_input_and_internal_failures_separately(
    monkeypatch, capsys, failure_type, kind
):
    """子进程区分无效输入和内部故障，只公开故障类别"""
    from tinkerfin_studio.attachments import worker

    def fail(_payload):
        raise failure_type("private generator details")

    monkeypatch.setattr(worker.resource, "setrlimit", lambda *_args: None)
    monkeypatch.setattr(
        worker.sys,
        "stdin",
        SimpleNamespace(
            buffer=io.BytesIO(b'{"operation":"generate","kind":"md","text":"ok"}')
        ),
    )
    monkeypatch.setattr(worker, "generate", fail)
    with pytest.raises(SystemExit) as exited:
        worker.main()
    assert exited.value.code == 1
    assert json.loads(capsys.readouterr().out) == {"kind": kind}


@pytest.mark.parametrize(
    "kind,error_type", [("invalid_input", ValueError), ("internal_error", RuntimeError)]
)
async def test_file_processor_preserves_worker_failure_category_and_reaps_process(
    monkeypatch, kind, error_type
):
    """无效参数与执行故障分别传播，已退出的处理进程仍会等待回收"""
    from tinkerfin_studio.attachments import documents

    process = SimpleNamespace(
        returncode=1,
        communicate=AsyncMock(return_value=(json.dumps({"kind": kind}).encode(), b"")),
        wait=AsyncMock(return_value=1),
    )
    monkeypatch.setattr(
        documents.asyncio, "create_subprocess_exec", AsyncMock(return_value=process)
    )
    with pytest.raises(error_type):
        await documents.DocumentProcessor().run(
            {"operation": "generate", "kind": "md", "text": "# report"}
        )
    process.wait.assert_awaited_once()


def test_attachment_tools_describe_all_model_visible_parameters(attachments):
    from tinkerfin_studio.attachments.tools import build_attachment_tools

    tools = build_attachment_tools(
        service=attachments,
        processor=attachments.documents,
        user_id=1,
        thread_id="schema-only",
        image_model=None,
    )
    for tool in tools:
        schema_type = tool.tool_call_schema
        assert isinstance(schema_type, type)
        assert issubclass(schema_type, BaseModel)
        schema = schema_type.model_json_schema()
        assert all(
            field.get("description") for field in schema.get("properties", {}).values()
        )


@pytest.mark.parametrize("extension", ["md", "markdown", "MD"])
async def test_markdown_upload_preserves_encoding_and_read_authorization(
    attachments, extension
):
    """Markdown 接受 BOM 并保留原件字节，下载仍要求附件所有权"""
    original = "\ufeff# 门店月报\r\n\r\n| 门店 | 营收 |\r\n| --- | --- |\r\n| 一店 | 128 |\r\n".encode()
    file = await attachments.upload(
        user_id=1, name=f"月报.{extension}", chunks=byte_chunks(original)
    )
    assert file.mime_type == "text/markdown"
    assert file.size_bytes == len(original)
    _, downloaded = await attachments.read(file.id, user_id=1)
    assert downloaded == original
    with pytest.raises(BusinessException):
        await attachments.read(file.id, user_id=2)
    with pytest.raises(BusinessException):
        await attachments.read(file.id, user_id=1, variant="preview")


@pytest.mark.parametrize("data", [b"", b"\xff\xfe#\x00", b"# report\x00data"])
async def test_invalid_markdown_is_not_published(attachments, database, data):
    """空文件、非 UTF-8 和二进制内容不得留下可见附件"""
    with pytest.raises(BusinessException):
        await attachments.upload(user_id=1, name="report.md", chunks=byte_chunks(data))
    async with database.session() as session:
        assert await session.scalar(select(AttachmentFile.id)) is None


@pytest.mark.parametrize("extension", ["md", "markdown"])
async def test_markdown_tools_generate_deliver_import_and_reopen(
    attachments, database, attachment_storage, extension, work_file_runtime
):
    """Markdown 生成、交付和重新导入保留原件，服务重建后仍校验会话权限"""
    import json

    from langchain_core.messages import ToolMessage

    from tinkerfin_contracts.media import attachment_from_block
    from tinkerfin_studio.attachments.tools import build_attachment_tools
    from tinkerfin_studio.attachments.workspace_tools import (
        build_sandbox_attachment_tools,
    )

    runtime, _, files = work_file_runtime

    async with database.session() as session:
        now = datetime.now(UTC).replace(tzinfo=None)
        session.add(
            ConversationThread(
                user_id=1,
                thread_id="markdown-report",
                title="门店月报",
                created_at=now,
                updated_at=now,
            )
        )
        await session.commit()
    tools = {
        item.name: item
        for item in build_attachment_tools(
            service=attachments,
            processor=attachments.documents,
            user_id=1,
            thread_id="markdown-report",
            image_model=None,
        )
    }
    text = "# 门店月报\n\n- 营收：128 万元\n- 下月安排：优化排班\n"
    result = await tools["create_file"].ainvoke(
        {
            "type": "tool_call",
            "id": "create-report",
            "name": "create_file",
            "args": {
                "name": f"月报.{extension}",
                "kind": "md",
                "text": text,
                "runtime": runtime,
            },
        }
    )
    assert isinstance(result, ToolMessage)
    assert isinstance(result.content, str)
    work_file = json.loads(result.content)
    assert files[work_file["file_path"]] == text.encode()
    assert work_file["shell_path"] == work_file["file_path"].lstrip("/")
    assert await attachments.list_thread(user_id=1, thread_id="markdown-report") == []
    deliver = build_sandbox_attachment_tools(
        service=attachments, user_id=1, thread_id="markdown-report"
    )[0]
    result = await deliver.ainvoke(
        {
            "type": "tool_call",
            "id": "deliver-report",
            "name": "deliver_file",
            "args": {
                "runtime": runtime,
                "file_path": work_file["file_path"],
                "name": work_file["name"],
            },
        }
    )
    assert isinstance(result, ToolMessage)
    assert isinstance(result.content, list)
    file = attachment_from_block(result.content[0])
    assert file is not None and file.mime_type == "text/markdown"
    restored = AttachmentService(database, attachment_storage)
    _, original = await restored.read(file.id, user_id=1, thread_id="markdown-report")
    assert original == text.encode()
    imported = json.loads(
        await tools["import_attachment"].ainvoke(
            {"attachment_id": file.id, "runtime": runtime}
        )
    )
    assert imported["file_path"] != work_file["file_path"]
    assert imported["shell_path"] == imported["file_path"].lstrip("/")
    assert files[imported["file_path"]] == original
    files[imported["file_path"]] = b"# edited"
    assert (await restored.read(file.id, user_id=1, thread_id="markdown-report"))[
        1
    ] == original
    assert [
        item.id
        for item in await restored.list_thread(user_id=1, thread_id="markdown-report")
    ] == [file.id]
    with pytest.raises(BusinessException):
        await restored.read(file.id, user_id=1, thread_id="another-thread")
    with pytest.raises(BusinessException):
        await restored.read(file.id, user_id=2, thread_id="markdown-report")
