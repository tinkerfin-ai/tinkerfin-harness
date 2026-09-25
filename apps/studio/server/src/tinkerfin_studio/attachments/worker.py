"""在可终止的独立进程内生成文件和处理图片预览"""

from __future__ import annotations

import base64
import io
import json
import math
import resource
import sys
from typing import Annotated, Literal
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from PIL import Image, ImageDraw, UnidentifiedImageError
from pydantic import BaseModel, ConfigDict, Field, JsonValue, TypeAdapter
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from tinkerfin_studio.attachments.processing import (
    MAX_FILE_BYTES,
    image_variant,
)


class _GenerateDocument(BaseModel):
    """子进程生成请求，只接收有界文字、单元格与图表数值"""

    model_config = ConfigDict(extra="forbid", allow_inf_nan=False)
    operation: Literal["generate"]
    kind: Literal["xlsx", "pdf", "png", "md"]
    text: str = Field(default="", max_length=50_000)
    rows: (
        list[Annotated[list[str | int | float | bool | None], Field(max_length=50)]]
        | None
    ) = Field(default=None, max_length=1000)
    values: list[Annotated[float, Field(ge=0, le=1e12)]] | None = Field(
        default=None, min_length=1, max_length=30
    )


class _ImagePreview(BaseModel):
    """为工作图片生成不超过模型读取限制的首帧预览"""

    model_config = ConfigDict(extra="forbid")
    operation: Literal["preview_image"]
    data: str = Field(max_length=14_000_000, description="原图片的 Base64 内容")
    max_bytes: int = Field(
        gt=0, le=MAX_FILE_BYTES, strict=True, description="预览允许的最大字节数"
    )


_DOCUMENT_REQUEST = TypeAdapter(
    Annotated[
        _GenerateDocument | _ImagePreview,
        Field(discriminator="operation"),
    ]
)


def preview_image(payload: _ImagePreview) -> dict[str, JsonValue]:
    """生成有界 JPEG 缩略图，不修改原图"""
    data = base64.b64decode(payload.data, validate=True)
    for max_side in (1280, 960, 640, 320):
        try:
            preview = image_variant(data, max_side)
        except OSError as error:
            # 图片从内存解码；此处的解码失败表示输入损坏，不是存储服务故障
            raise ValueError("图片内容无法解码") from error
        if len(preview) <= payload.max_bytes:
            return {"data": base64.b64encode(preview).decode("ascii")}
    raise ValueError("图片预览超过读取限制")


def generate(payload: _GenerateDocument) -> dict[str, JsonValue]:
    """生成 Markdown、XLSX、PDF 或 PNG，限制输入规模"""
    kind = payload.kind
    output = io.BytesIO()
    if kind == "xlsx":
        rows = payload.rows
        if not isinstance(rows, list) or len(rows) > 1000:
            raise ValueError("表格最多 1000 行")
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet("数据")
        for row in rows:
            if not isinstance(row, list) or len(row) > 50:
                raise ValueError("表格最多 50 列")
            cells = []
            for value in row:
                if isinstance(value, float) and not math.isfinite(value):
                    raise ValueError("表格数值必须为有限数")
                cell = WriteOnlyCell(
                    sheet, value=value[:2000] if isinstance(value, str) else value
                )
                if isinstance(value, str):
                    cell.data_type = "s"
                cells.append(cell)
            sheet.append(cells)
        workbook.save(output)
    elif kind == "md":
        output.write(payload.text.encode("utf-8"))
    elif kind == "pdf":
        text = payload.text
        if len(text) > 50_000:
            raise ValueError("报告文字超过限制")
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        style = ParagraphStyle(
            "report", fontName="STSong-Light", fontSize=11, leading=18, wordWrap="CJK"
        )
        content = []
        for line in text.splitlines():
            content.extend([Paragraph(escape(line) or " ", style), Spacer(1, 6)])
        SimpleDocTemplate(output).build(content)
    elif kind == "png":
        values = payload.values
        if (
            not isinstance(values, list)
            or not 1 <= len(values) <= 30
            or not all(isinstance(v, (int, float)) and 0 <= v <= 1e12 for v in values)
        ):
            raise ValueError("图表需要 1 到 30 个非负数值")
        image = Image.new("RGB", (1000, 600), "white")
        draw = ImageDraw.Draw(image)
        draw.line((60, 40, 60, 540, 960, 540), fill="#61666b", width=2)
        maximum = max(values) or 1
        width = 850 / len(values)
        for index, value in enumerate(values):
            left = 80 + index * width
            top = 530 - 440 * value / maximum
            draw.rectangle((left, top, left + width * 0.7, 538), fill="#3158e6")
            draw.text((left, top - 18), str(value), fill="#111111")
            draw.text((left, 552), str(index + 1), fill="#111111")
        image.save(output, "PNG")
    else:
        raise ValueError("生成类型只支持 Markdown、XLSX、PDF 或 PNG")
    data = output.getvalue()
    if len(data) > 10 * 1024 * 1024:
        raise ValueError("生成文件超过 10 MiB")
    return {"data": base64.b64encode(data).decode("ascii")}


def main() -> None:
    """执行一个标准输入任务并输出有界 JSON，不读取用户任意路径"""
    try:
        resource.setrlimit(resource.RLIMIT_CPU, (25, 26))
        if sys.platform.startswith("linux"):
            resource.setrlimit(
                resource.RLIMIT_AS, (768 * 1024 * 1024, 768 * 1024 * 1024)
            )
        payload = _DOCUMENT_REQUEST.validate_json(
            sys.stdin.buffer.read(20 * 1024 * 1024)
        )
        if isinstance(payload, _ImagePreview):
            result = preview_image(payload)
        else:
            result = generate(payload)
        encoded = json.dumps(result, ensure_ascii=False)
        if len(encoded) > 15_000_000:
            raise ValueError("生成结果超过大小限制")
        print(encoded)
    except (
        ValueError,
        UnidentifiedImageError,
        Image.DecompressionBombWarning,
        Image.DecompressionBombError,
    ):
        print(
            json.dumps(
                {"kind": "invalid_input"},
                ensure_ascii=False,
            )
        )
        raise SystemExit(1)
    except Exception:  # noqa: BLE001 - 子进程返回故障类别，父进程继续传播运行失败
        print(json.dumps({"kind": "internal_error"}))
        raise SystemExit(1)


if __name__ == "__main__":
    main()
